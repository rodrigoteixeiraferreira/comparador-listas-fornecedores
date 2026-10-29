"""O gerador de corrupções: estraga uma lista boa de um jeito conhecido, para provar que o portão segura.

Cada corrupção diz em que conferência do portão ela tem de cair. A calibração (tests/test_gate.py) roda todas, em
várias sementes, e confere que cada uma cai na conferência certa e que nenhuma lista boa é segurada.
"""
from __future__ import annotations

import copy
import random
from dataclasses import dataclass, replace
from pathlib import Path

from openpyxl import load_workbook

from ..reference import Reference, load
from .build import VAT, PriceList
from .files import write_list


@dataclass(frozen=True)
class Corruption:
    code: str
    formats: tuple[str, ...]  # em que formato de arquivo ela faz sentido
    expected_check: int  # a conferência do portão que tem de segurar
    what: str


CORRUPTIONS = (
    Corruption("header_renamed", ("xlsx",), 1, "o fornecedor trocou o rótulo da coluna da descrição"),
    Corruption("price_format_changed", ("pdf",), 1, "parte dos preços veio em outro formato ($4850.00)"),
    Corruption("price_column_moved", ("xlsx",), 2, "os rótulos de preço e de unidades por caixa trocaram de lugar"),
    Corruption("truncated", ("xlsx", "pdf"), 3, "o arquivo chegou cortado pela metade"),
    Corruption("codes_renumbered", ("xlsx", "pdf"), 3, "o fornecedor renumerou todos os códigos"),
    Corruption("prices_x10", ("xlsx", "pdf"), 5, "todos os preços com um zero a mais"),
    Corruption("vat_flipped", ("xlsx", "pdf"), 5, "o IVA declarado não é o dos preços"),
)


def corrupt(price_list: PriceList, kind: str, folder: Path, seed: int = 0, ref: Reference | None = None) -> Path:
    """Grava em `folder` a lista estragada do jeito `kind` e devolve o caminho."""
    ref = ref or load()
    rng = random.Random(f"{seed}:{price_list.name}:{kind}")
    pl = copy.deepcopy(price_list)
    price_text = None
    if kind == "truncated":
        pl.lines = pl.lines[: int(len(pl.lines) * 0.45)]
    elif kind == "codes_renumbered":
        prefix = "" if pl.supplier == "mayorista_norte" else "B"
        codes = rng.sample(range(10000, 99999), len(pl.lines))
        pl.lines = [replace(ln, supplier_code=f"{prefix}{c}") for ln, c in zip(pl.lines, codes)]
    elif kind == "prices_x10":
        pl.lines = [replace(ln, price=round(ln.price * 10, 2)) for ln in pl.lines]
    elif kind == "vat_flipped":
        # a lista diz «com IVA» e manda sem, ou diz «sem IVA» e manda com
        factor = 1 / VAT if pl.prices_include_vat else VAT
        pl.lines = [replace(ln, price=round(ln.price * factor, 2)) for ln in pl.lines]
    elif kind == "price_format_changed":
        changed = set(rng.sample(range(len(pl.lines)), int(len(pl.lines) * 0.4)))
        rows = {id(ln): i for i, ln in enumerate(pl.lines)}

        def price_text(line):
            if rows[id(line)] in changed:
                return f"${line.price:.2f}"
            return "$ " + f"{line.price:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    elif kind not in ("header_renamed", "price_column_moved"):
        raise ValueError(f"corrupção desconhecida: {kind}")
    path = write_list(pl, folder, ref, price_text=price_text)
    if kind in ("header_renamed", "price_column_moved"):
        wb = load_workbook(path)
        ws = wb.active
        header = next(r for r in ws.iter_rows() if any(c.value == "DESCRIPCION" for c in r))
        cells = {c.value: c for c in header}
        if kind == "header_renamed":
            cells["DESCRIPCION"].value = "DESC. DEL PRODUCTO"
        else:
            a, b = cells["UNID. X BULTO"], cells["PRECIO S/IVA"]
            a.value, b.value = b.value, a.value
        wb.save(path)
    return path
