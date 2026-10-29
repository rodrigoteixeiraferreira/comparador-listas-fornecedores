"""Grava o mundo em disco: o cadastro da loja, as listas como os fornecedores mandam e o gabarito, à parte.

    <pasta>/catalog.csv              o cadastro da loja (o que o sistema recebe)
    <pasta>/lists/<fornecedor>_<mês>.xlsx|.pdf
    <pasta>/answer_key/catalog.csv   o que cada produto do cadastro é de verdade
    <pasta>/answer_key/lines.csv     o que cada linha de lista é de verdade, e a armadilha que carrega
"""
from __future__ import annotations

import csv
from collections.abc import Callable
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font

from ..reference import Reference, load
from .build import ListLine, PriceList, World
from .spec import MONTH_NAMES_ES


def write_world(world: World, out: Path, ref: Reference | None = None) -> list[Path]:
    ref = ref or load()
    (out / "lists").mkdir(parents=True, exist_ok=True)
    (out / "answer_key").mkdir(parents=True, exist_ok=True)
    written = [_write_catalog(world, out / "catalog.csv")]
    for price_list in world.lists:
        written.append(write_list(price_list, out / "lists", ref))
    written += _write_answer_key(world, out / "answer_key")
    return written


def write_list(price_list: PriceList, folder: Path, ref: Reference | None = None,
               price_text: Callable[[ListLine], str] | None = None) -> Path:
    """Uma lista no formato do fornecedor. `price_text` muda como o preço sai no PDF (para as corrupções)."""
    from .photo import write_photo

    ref = ref or load()
    supplier = ref.suppliers[price_list.supplier]
    path = folder / f"{price_list.name}.{supplier.extension}"
    folder.mkdir(parents=True, exist_ok=True)
    if supplier.file_format == "xlsx":
        _write_xlsx(price_list, path, ref)
    elif supplier.file_format == "image":
        write_photo(price_list, path, ref)
    else:
        _write_pdf(price_list, path, ref, price_text or (lambda line: _ars(line.price)))
    return path


def _write_catalog(world: World, path: Path) -> Path:
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["product_id", "barcode", "description", "cost"])
        for c in world.catalog:
            w.writerow([c.id, c.barcode, c.description, f"{c.cost:.2f}"])
    return path


def _write_answer_key(world: World, folder: Path) -> list[Path]:
    market = world.market_by_id()
    catalog_path = folder / "catalog.csv"
    with open(catalog_path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["product_id", "market_id", "type", "brand", "variety", "size", "pack"])
        for c in world.catalog:
            p = market[c.market_id]
            w.writerow([c.id, p.id, p.type_code, p.brand, p.variety, p.size, p.pack])
    lines_path = folder / "lines.csv"
    with open(lines_path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["list", "row", "supplier_code", "description", "market_id", "catalog_id", "type", "brand",
                    "variety", "size", "pack", "extra", "traps", "units_per_box", "price_per_box", "printed_price",
                    "true_unit_price", "price_error"])
        for pl in world.lists:
            for line in pl.lines:
                p = market[line.market_id]
                w.writerow([pl.name, line.row, line.supplier_code, line.description, p.id, line.catalog_id or "",
                            p.type_code, p.brand, p.variety, p.size, p.pack, p.extra, ";".join(line.traps),
                            line.units_per_box, str(line.price_per_box).lower(), f"{line.price:.2f}",
                            f"{line.true_unit_price:.2f}", f"{line.price_error:g}"])
    return [catalog_path, lines_path]


def _write_xlsx(price_list: PriceList, path: Path, ref: Reference) -> None:
    """A Mayorista Norte manda planilha: título em cima, preço sem IVA, a caixa numa coluna à parte."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.append([ref.suppliers[price_list.supplier].name.upper()])
    ws.append([f"LISTA DE PRECIOS - {MONTH_NAMES_ES[price_list.period]}"])
    ws.append(["PRECIOS SIN IVA. VENTA POR BULTO CERRADO."])
    ws.append([])
    ws.append(["RUBRO", "CODIGO", "EAN", "DESCRIPCION", "UNID. X BULTO", "PRECIO S/IVA"])
    for cell in ws[1] + ws[5]:
        cell.font = Font(bold=True)
    for line in price_list.lines:
        ws.append([ref.sections[line.section].name_es, int(line.supplier_code), line.barcode or None,
                   line.description, line.units_per_box, line.price])
        ws.cell(ws.max_row, 6).number_format = "#,##0.00"
    for col, width in zip("ABCDEF", (12, 9, 15, 48, 14, 14)):
        ws.column_dimensions[col].width = width
    wb.save(path)


def _ars(value: float) -> str:
    """Preço no formato argentino: $ 4.850,00."""
    return "$ " + f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


class _AndinaPdf(FPDF):
    def __init__(self, title: str, subtitle: str):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.title_text = title
        self.subtitle = subtitle
        self.set_auto_page_break(auto=True, margin=15)

    def header(self) -> None:
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 7, self.title_text, new_x="LMARGIN", new_y="NEXT")
        self.set_font("Helvetica", "", 9)
        self.cell(0, 5, self.subtitle, new_x="LMARGIN", new_y="NEXT")
        self.ln(2)
        self.set_font("Helvetica", "B", 9)
        self.cell(22, 6, "Cod.", border="B")
        self.cell(128, 6, "Articulo", border="B")
        self.cell(30, 6, "Precio", border="B", align="R", new_x="LMARGIN", new_y="NEXT")

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 5, f"Pagina {self.page_no()}", align="C")


def _write_pdf(price_list: PriceList, path: Path, ref: Reference, price_text: Callable[[ListLine], str]) -> None:
    """A Distribuidora Andina manda PDF: preço com IVA, por seção; a linha com CAJA tem o preço da caixa."""
    name = ref.suppliers[price_list.supplier].name.upper()
    pdf = _AndinaPdf(f"{name} - LISTA DE PRECIOS {MONTH_NAMES_ES[price_list.period]}",
                     "Precios con IVA incluido. Los articulos con CAJA se venden por caja cerrada.")
    pdf.add_page()
    section = None
    for line in price_list.lines:
        if line.section != section:
            section = line.section
            pdf.ln(1)
            pdf.set_font("Helvetica", "B", 9)
            pdf.cell(0, 6, ref.sections[section].name_es, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        pdf.cell(22, 5, line.supplier_code)
        pdf.cell(128, 5, line.description)
        pdf.cell(30, 5, price_text(line), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))
