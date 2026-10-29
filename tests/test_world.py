"""O gerador do mundo fictício: determinístico, com todas as armadilhas e com um gabarito coerente."""
import pdfplumber
import pytest
from openpyxl import load_workbook

from pricelists.world import build_world, write_world
from pricelists.world import build as b

SEEDS = range(1, 6)
ALL_TRAPS = {b.SIZE_SIBLING, b.BRAND_SWAP, b.VARIETY_CONFLICT, b.PACK_VS_SINGLE, b.LEFTOVER_WORD,
             b.UNKNOWN_WRITING, b.NOISE_WORD, b.BOX_PRICE, b.STALE_BARCODE, b.REUSED_CODE, b.PRICE_FACTOR}


def _snapshot(world):
    return [(pl.name, ln.supplier_code, ln.description, ln.price, ln.traps) for pl in world.lists for ln in pl.lines]


def test_same_seed_same_world():
    assert _snapshot(build_world(1)) == _snapshot(build_world(1))
    assert _snapshot(build_world(1)) != _snapshot(build_world(2))


@pytest.mark.parametrize("seed", SEEDS)
def test_barcodes_are_internal_range_and_valid(seed):
    world = build_world(seed)
    codes = [p.barcode for p in world.market]
    assert len(codes) == len(set(codes))
    for code in codes:
        assert len(code) == 13 and code[0] == "2"  # faixa de uso interno: nenhum produto real
        total = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(code[:12]))
        assert int(code[12]) == (10 - total % 10) % 10


@pytest.mark.parametrize("seed", SEEDS)
def test_every_trap_is_planted(seed):
    world = build_world(seed)
    planted = {t for pl in world.lists for ln in pl.lines for t in ln.traps}
    assert planted == ALL_TRAPS


@pytest.mark.parametrize("seed", SEEDS)
def test_catalog_size_and_unique_descriptions(seed):
    world = build_world(seed)
    assert 150 <= len(world.catalog) <= 300
    descriptions = [c.description for c in world.catalog]
    assert len(descriptions) == len(set(descriptions))


@pytest.mark.parametrize("seed", SEEDS)
def test_answer_key_is_consistent(seed):
    world = build_world(seed)
    market = world.market_by_id()
    catalog = {c.id: c for c in world.catalog}
    catalog_by_barcode = {c.barcode: c for c in world.catalog if c.barcode}
    for pl in world.lists:
        codes = [ln.supplier_code for ln in pl.lines]
        assert len(codes) == len(set(codes)), "código repetido na mesma lista"
        for ln in pl.lines:
            p = market[ln.market_id]
            if ln.catalog_id is not None:
                assert catalog[ln.catalog_id].market_id == p.id
            if p.extra:
                assert ln.catalog_id is None  # a loja nunca tem o produto com a palavra a mais
            if b.STALE_BARCODE in ln.traps:
                other = catalog_by_barcode[ln.barcode]
                assert other.market_id != p.id and market[other.market_id].size != p.size
            # o preço impresso, desfeito o IVA, a caixa e o erro plantado, volta ao preço certo
            price = ln.price / ln.price_error
            if not pl.prices_include_vat:
                price *= b.VAT
            if ln.price_per_box:
                price /= ln.units_per_box * 0.97
            assert price == pytest.approx(ln.true_unit_price, abs=0.02)
            assert (ln.price_error != 1.0) == (b.PRICE_FACTOR in ln.traps)


def test_codes_carry_over_between_months():
    world = build_world(1)
    for supplier in ("mayorista_norte", "distribuidora_andina"):
        sep, oct_ = [pl for pl in world.lists if pl.supplier == supplier]
        sep_codes = {ln.supplier_code: ln.market_id for ln in sep.lines}
        kept = [ln for ln in oct_.lines if ln.supplier_code in sep_codes]
        reused = [ln for ln in kept if sep_codes[ln.supplier_code] != ln.market_id]
        assert len(kept) > 0.8 * len(sep.lines)
        assert reused and all(b.REUSED_CODE in ln.traps for ln in reused)


def test_files_carry_every_line(tmp_path):
    world = build_world(1)
    write_world(world, tmp_path)
    assert (tmp_path / "catalog.csv").exists()
    assert (tmp_path / "answer_key" / "lines.csv").exists()
    for pl in world.lists:
        if pl.supplier == "mayorista_norte":
            ws = load_workbook(tmp_path / "lists" / f"{pl.name}.xlsx").active
            rows = list(ws.iter_rows(min_row=6, values_only=True))
            assert [str(r[1]) for r in rows] == [ln.supplier_code for ln in pl.lines]
        else:
            with pdfplumber.open(tmp_path / "lists" / f"{pl.name}.pdf") as pdf:
                text = "\n".join(page.extract_text() for page in pdf.pages)
            assert all(ln.supplier_code in text for ln in pl.lines)
